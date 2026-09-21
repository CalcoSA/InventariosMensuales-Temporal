from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import re
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from app.constants import BASES_FOLDER_NAME, FOLDER_MIME, SHEET_MIME
from app.models import generator_rules as rules
from app.models.errors import DomainError
from app.models.sheets import find_sheet


def build_xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BASE"
    sheet.append(["ÍTEM", "PRODUCTO", "U.D MED", "CATEGORÍA"])
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="center")
            if sheet.max_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3E9DC")
        sheet.cell(sheet.max_row, 1).number_format = "@"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="6B3F2A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for column, pixels in zip("ABCD", (110, 390, 180, 190)):
        sheet.column_dimensions[column].width = (pixels - 5) / 7
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:D{len(rows) + 1}"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class MonthlyGeneratorService:
    def __init__(self, repository, locks, state_file, monthly_folder, formats_folder, timezone="America/Bogota", now=None):
        self.repository, self.locks = repository, locks
        self.state_file = Path(state_file)
        self.monthly_folder, self.formats_folder = monthly_folder, formats_folder
        self.now = now or (lambda: datetime.now(ZoneInfo(timezone)))

    def status(self):
        if not self.state_file.exists():
            return {"ESTADO": "NO INICIADO"}
        return json.loads(self.state_file.read_text(encoding="utf-8"))

    def _save(self, state):
        self.state_file.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.state_file.with_suffix(".tmp")
        temporary.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(self.state_file)

    def reset(self):
        with self.locks.hold("monthly-generator"):
            if self.state_file.exists():
                self.state_file.unlink()

    def process_pdv(self, monthly, templates, folder):
        template = rules.best_template(monthly["nombre"], templates)
        if not template:
            raise DomainError("No se encontró un formato correspondiente.")
        with self.repository.open_book(monthly) as book:
            with self.repository.open_book(template) as format_book:
                products = rules.extract_products(book)
                if not products:
                    raise DomainError("No se encontraron productos en Pedido Diario o Bodega.")
                categories = rules.categories_by_code(format_book, products)
                if not any(categories.values()):
                    raise DomainError("No se detectó ninguna categoría dentro de la hoja Mensual.")
                return self.repository.write_xlsx(folder, rules.output_name(monthly["nombre"]),
                                                  build_xlsx(rules.output_rows(products, categories)))

    def run(self, *, limit=None, progress=None):
        with self.locks.hold("monthly-generator"):
            state = self.status()
            if state["ESTADO"] == "COMPLETADO":
                return state
            if state["ESTADO"] == "NO INICIADO":
                monthly = [f for f in self.repository.files(self.monthly_folder) if re.search(r"09\s*\.xlsx$", f["nombre"], re.I)]
                templates = rules.deduplicate_templates(self.repository.files(self.formats_folder))
                if not monthly:
                    raise DomainError("No se encontraron archivos .xlsx terminados en 09.")
                if not templates:
                    raise DomainError("No se encontraron archivos en la carpeta de formatos.")
                names = [rules.output_name(f["nombre"]) for f in monthly]
                if len(set(names)) != len(names):
                    raise DomainError("Dos archivos mensuales producirían el mismo nombre de salida. Revise los archivos origen.")
                state = dict(ESTADO="CREANDO CARPETA", INDICE=0, TOTAL=len(monthly), ARCHIVOS_MENSUALES=monthly,
                             ARCHIVOS_FORMATOS=templates, ERRORES=[], INICIO=self.now().isoformat(),
                             NOMBRE_SALIDA="RESULTADOS INVENTARIO MENSUAL " + self.now().strftime("%Y-%m-%d %H-%M"))
                # A reset within the same minute starts a new job, not a reuse of old results.
                # Record pre-existing IDs BEFORE creating, so an ambiguous create can resume safely.
                state["CARPETAS_PREEXISTENTES"] = [f["id"] for f in self.repository.drive.list(
                    self.monthly_folder, name=state["NOMBRE_SALIDA"], mime=FOLDER_MIME)]
                self._save(state)
            if not state.get("CARPETA_SALIDA_ID"):
                folders = self.repository.drive.list(self.monthly_folder, name=state["NOMBRE_SALIDA"], mime=FOLDER_MIME)
                folders = [f for f in folders if f["id"] not in state.get("CARPETAS_PREEXISTENTES", [])]
                if len(folders) > 1:
                    raise DomainError("Hay varias carpetas con el nombre de salida. Revise antes de continuar.")
                folder = folders[0] if folders else self.repository.drive.create_folder(self.monthly_folder, state["NOMBRE_SALIDA"])
                state["CARPETA_SALIDA_ID"] = folder["id"]
                self._save(state)
            done = 0
            while state["INDICE"] < state["TOTAL"] and (limit is None or done < limit):
                monthly = state["ARCHIVOS_MENSUALES"][state["INDICE"]]
                state["ESTADO"] = "EN PROCESO"
                self._save(state)
                try:
                    # A prior interrupted upload is discovered before repeating conversion.
                    name = rules.output_name(monthly["nombre"])
                    existing = self.repository.drive.list(state["CARPETA_SALIDA_ID"], name=name)
                    if len(existing) > 1:
                        raise DomainError("Existen resultados duplicados para " + name + ".")
                    if not existing:
                        self.process_pdv(monthly, state["ARCHIVOS_FORMATOS"], state["CARPETA_SALIDA_ID"])
                except DomainError as error:
                    state["ERRORES"].append(monthly["nombre"] + ": " + str(error))
                state["INDICE"] += 1
                done += 1
                state["ULTIMA_ACTUALIZACION"] = self.now().isoformat()
                self._save(state)
                if progress:
                    progress(state)
            if state["INDICE"] < state["TOTAL"]:
                state["ESTADO"] = "PENDIENTE DE CONTINUAR"
            else:
                state["ESTADO"] = "CREANDO ZIP"
                self._save(state)
                zip_file = self.repository.zip_outputs(state["CARPETA_SALIDA_ID"])
                state.update(ESTADO="COMPLETADO", ZIP_ID=zip_file["id"],
                             ZIP_URL=zip_file.get("webViewLink", "https://drive.google.com/file/d/" + zip_file["id"] + "/view"),
                             FIN=self.now().isoformat())
            self._save(state)
            return state


class MonthlyPreparationService:
    def __init__(self, bases, locks):
        self.bases, self.locks = bases, locks

    def run(self, limit=5):
        if not isinstance(limit, int) or limit < 1:
            raise DomainError("El límite debe ser un entero positivo.")
        drive = self.bases.drive
        with self.locks.hold("monthly-preparation"):
            folders = drive.list(self.bases.formats_folder, name=BASES_FOLDER_NAME, mime=FOLDER_MIME)
            folder = folders[0] if folders else drive.create_folder(self.bases.formats_folder, BASES_FOLDER_NAME)
            existing = {f["name"] for f in drive.list(folder["id"], mime=SHEET_MIME)}
            files = [f for f in drive.list(self.bases.formats_folder) if f["name"].lower().endswith(".xlsx")]
            converted = processed = attempts = 0
            errors = []
            for file in files:
                name = re.sub(r"\.xlsx$", "", file["name"], flags=re.I).strip()
                if name in existing:
                    converted += 1
                    continue
                if attempts >= limit:
                    continue
                attempts += 1
                try:
                    created = drive.upload(folder["id"], name, drive.download(file["id"]), file["mimeType"], convert=True)
                    book = self.bases.sheets.read_book(created["id"])
                    if not find_sheet(book, "Mensual"):
                        drive.trash(created["id"])
                        raise DomainError("El archivo no contiene la hoja Mensual.")
                    existing.add(name)
                    converted += 1
                    processed += 1
                    self.bases.cache.clear()
                except DomainError as error:
                    errors.append(file["name"] + ": " + str(error))
            return dict(estado="COMPLETADO" if converted == len(files) else "PENDIENTE DE CONTINUAR",
                        convertidos=converted, total=len(files), procesadosAhora=processed, pendientes=len(files) - converted,
                        carpeta="https://drive.google.com/drive/folders/" + folder["id"], errores=errors)
