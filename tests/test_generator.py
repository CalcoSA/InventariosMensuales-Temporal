from copy import deepcopy
from io import BytesIO
import json
import zipfile
from openpyxl import load_workbook
import pytest
from app.constants import SHEET_MIME, XLSX_MIME
from app.models.errors import DomainError
from app.services.monthly_generator import build_xlsx
from tests.fakes import sheet


def seed_generator(c):
    c.drive.files.extend([
        dict(id="monthly-xlsx",name="BR01 - Norte 09.xlsx",mimeType=XLSX_MIME,parents=["monthly"]),
        dict(id="format-sheet",name="BR01 - Norte.xlsx",mimeType=SHEET_MIME,parents=["formats"]),
        dict(id="ignore",name="BR01 - Norte 08.xlsx",mimeType=XLSX_MIME,parents=["monthly"]),
    ])
    c.drive.content["monthly-xlsx"]=b"fake xlsx"
    c.sheets.books["format-sheet"]=[sheet("Mensual",[
        ["Código","Categoría"],["00123","Bebidas"],["00456","Cocina"],
    ])]
    c.drive.on_convert=lambda result,content:c.sheets.books.update({result["id"]:[
        sheet("Pedido Diario",[["Código","Producto","Unidad"],["00123","Jugo X 750 ML","NO USAR"],
                               ["00456","Queso X 2 KG","NO USAR"]]),
        sheet("Bodega",[["Código","Producto"],["00123","Nombre secundario"],["00789","Azúcar X 12"]],1)
    ]})


def test_real_xlsx_format_and_text():
    data=build_xlsx([["00123","Jugo X 750 ML","X 750 ML","Bebidas"],["=9","=HYPERLINK(\"x\")","","Cocina"]])
    book=load_workbook(BytesIO(data))
    ws=book["BASE"]
    assert list(next(ws.values))==["ÍTEM","PRODUCTO","U.D MED","CATEGORÍA"]
    assert ws["A2"].value=="00123" and ws["A2"].number_format=="@"
    assert ws["A3"].data_type=="s" and ws["B3"].data_type=="s"
    assert ws["A1"].fill.fgColor.rgb.endswith("6B3F2A")
    assert ws["A1"].font.bold and ws["A1"].font.color.rgb.endswith("FFFFFF")
    assert ws["A2"].fill.fgColor.rgb.endswith("F3E9DC")
    assert ws.freeze_panes=="A2" and ws.auto_filter.ref=="A1:D3"
    book.close()


def test_generator_full_run_resume_zip_and_temporaries(container):
    seed_generator(container)
    state=container.generator.run()
    assert state["ESTADO"]=="COMPLETADO" and state["TOTAL"]==1
    assert state["ERRORES"]==[]
    folder=state["CARPETA_SALIDA_ID"]
    files=[f for f in container.drive.files if folder in f["parents"]]
    assert sorted(f["name"] for f in files)==["BR01 - Norte - BASE INVENTARIO.xlsx","INVENTARIOS_PDV_GENERADOS.zip"]
    workbook=load_workbook(BytesIO(container.drive.content[next(f["id"] for f in files if f["name"].endswith(".xlsx"))]))
    values=list(workbook["BASE"].values)
    assert any(r[0]=="00123" and r[1]=="Jugo X 750 ML" and r[2]=="X 750 ML" for r in values)
    assert any(r[0]=="00789" and r[3]=="SIN CATEGORÍA" for r in values)
    with zipfile.ZipFile(BytesIO(container.drive.content[state["ZIP_ID"]])) as archive:
        assert archive.namelist()==["BR01 - Norte - BASE INVENTARIO.xlsx"]
    assert all(f.get("trashed") for f in container.drive.files if f["name"].startswith("TEMPORAL_"))
    writes=container.drive.writes
    assert container.generator.run()["ZIP_ID"]==state["ZIP_ID"]
    assert container.drive.writes==writes
    workbook.close()


def test_generator_interruption_after_upload_resumes_without_duplicate(container,monkeypatch):
    seed_generator(container)
    original=container.generator.process_pdv
    def interrupted(*args):
        original(*args)
        raise KeyboardInterrupt()
    monkeypatch.setattr(container.generator,"process_pdv",interrupted)
    with pytest.raises(KeyboardInterrupt):container.generator.run()
    monkeypatch.setattr(container.generator,"process_pdv",original)
    result=container.generator.run()
    folder=result["CARPETA_SALIDA_ID"]
    assert len([f for f in container.drive.files if folder in f["parents"] and f["name"].endswith(".xlsx")])==1
    assert result["ESTADO"]=="COMPLETADO"


def test_no_monthly_sources_does_not_create_folder(container):
    with pytest.raises(DomainError,match="No se encontraron archivos .xlsx terminados en 09."):
        container.generator.run()
    assert container.drive.writes==0


def test_explicit_reset_same_minute_does_not_reuse_stale_results(container):
    from datetime import datetime
    from zoneinfo import ZoneInfo
    seed_generator(container)
    container.generator.now=lambda:datetime(2026,9,18,12,0,tzinfo=ZoneInfo("America/Bogota"))
    first=container.generator.run()
    container.generator.reset()
    second=container.generator.run()
    assert first["CARPETA_SALIDA_ID"]!=second["CARPETA_SALIDA_ID"]
    assert first["ZIP_ID"]!=second["ZIP_ID"]


def test_preparation_limit_existing_and_invalid(container):
    for i in range(8):
        container.drive.files.append(dict(id=f"xlsx-{i}",name=f"PDV-{i}.xlsx",mimeType=XLSX_MIME,parents=["formats"]))
        container.drive.content[f"xlsx-{i}"]=str(i).encode()
    container.drive.on_convert=lambda result,content:container.sheets.books.update({
        result["id"]:[sheet("Otra" if content==b"0" else "Mensual",[["Categoría","Item","Producto"],["A","1","P"]])]
    })
    result=container.preparation.run()
    assert result["procesadosAhora"]==4
    assert result["pendientes"]==4 and len(result["errores"])==1
    created=[f for f in container.drive.files if f["name"]=="PDV-0"]
    assert created[0]["trashed"]
    result=container.preparation.run()
    assert result["convertidos"]==7 and result["pendientes"]==1
    assert len([f for f in container.drive.files if f["name"]=="PDV-1"])==1
